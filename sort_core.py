# -*- coding: utf-8 -*-
"""
ソート処理コアエンジン
Excelファイルの読み込み、整形、並べ替え、書式設定を行います。
"""

import sys
from pathlib import Path
from typing import Any, Callable, Dict, Optional, Tuple
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Border, Side, Alignment, Font

from config_manager import load_config

# Excelの列見出し（標準化後）
COLS_LOAN = ['融資実行日', '形態', 'お客様氏名', '物件', '依頼内容', '担当', '管轄', '立会時間', '立会場所', '立会者', '当日申請']
COLS_CANCEL = ['金消日・面談日', '形態', 'お客様氏名', '物件', '依頼内容', '担当', '管轄', '金消時間', '金消場所・面談場所', '意思確認', '融資実行日']
COLS_OUT_ORDER = ['日付', '形態', 'お客様氏名', '物件', '依頼内容', '担当', '管轄', '時間', '場所', '確認者', '申請', '識別']

def _read_sheet(path: Path, log_cb: Optional[Callable[[str], None]] = None) -> pd.DataFrame:
    """指定されたExcelファイルから対象シートを読み込みます"""
    if log_cb:
        log_cb(f"ファイルを読み込み中: {path.name}")
    excel_file = pd.ExcelFile(path)
    sheets = excel_file.sheet_names

    target_sheet = None
    candidates = ['受諾確認票', '受託確認票', 'Sheet1', 'シート1']
    for candidate in candidates:
        if candidate in sheets:
            target_sheet = candidate
            break

    if target_sheet is None and len(sheets) > 0:
        target_sheet = sheets[0]

    if log_cb:
        log_cb(f"対象シート: {target_sheet}")

    df = pd.read_excel(excel_file, sheet_name=target_sheet)
    df.columns = df.columns.astype(str).str.strip()

    # 日付列のパース
    for c in ['融資実行日', '金消日・面談日']:
        if c in df.columns:
            df[c] = pd.to_datetime(df[c], errors='coerce')
    return df

def _dt_to_str(s: pd.Series) -> pd.Series:
    """日付型Seriesを YYYY/MM/DD 形式の文字列へ変換します"""
    return s.dt.strftime('%Y/%m/%d')

def _parse_time_like(s: pd.Series) -> pd.Series:
    """
    時間文字列（HH:MM、H:MM、HH:MM:SSなど）を秒数に変換します。
    """
    if s is None:
        return pd.Series(dtype='float64')
    ss = s.fillna('')
    def tosec(x):
        t = str(x).strip()
        if not t:
            return np.nan
        parts = t.split(':')
        try:
            h = int(parts[0])
            m = int(parts[1]) if len(parts) > 1 else 0
            sec = int(parts[2]) if len(parts) > 2 else 0
            return h * 3600 + m * 60 + sec
        except Exception:
            return np.nan
    return ss.map(tosec).astype('float64')

def _make_tables(df: pd.DataFrame) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """融資実行日データ、金消日データ、統合データを生成します"""
    # 融資実行日データ
    loan = df.copy()
    if '融資実行日' in loan.columns:
        loan = loan.sort_values(by='融資実行日')
    loan = loan.assign(識別='融資実行日')
    loan = loan.rename(columns={
        '融資実行日': '日付',
        '立会時間': '時間',
        '立会場所': '場所',
        '立会者': '確認者',
        '当日申請': '申請'
    })
    loan = loan[[c for c in COLS_OUT_ORDER if c in loan.columns]]
    if '日付' in loan.columns:
        loan['日付'] = pd.to_datetime(loan['日付'], errors='coerce')
        loan['日付'] = _dt_to_str(loan['日付'])

    # 金消日・面談日データ
    cancel = df.copy()
    if '金消日・面談日' in cancel.columns:
        cancel = cancel.sort_values(by='金消日・面談日')
    cancel = cancel.assign(識別='金消日・面談日')
    cancel = cancel.rename(columns={
        '金消日・面談日': '日付',
        '金消時間': '時間',
        '金消場所・面談場所': '場所',
        '意思確認': '確認者',
        '融資実行日': '申請'
    })
    cancel = cancel[[c for c in COLS_OUT_ORDER if c in cancel.columns]]
    if '日付' in cancel.columns:
        cancel['日付'] = pd.to_datetime(cancel['日付'], errors='coerce')
        cancel['日付'] = _dt_to_str(cancel['日付'])

    # 統合データ
    combined = pd.concat([loan, cancel], ignore_index=True, sort=False)
    if '日付' in combined.columns:
        tmp = pd.to_datetime(combined['日付'], errors='coerce')
        combined = combined.assign(_d=tmp).sort_values(by=['_d']).drop(columns=['_d'])
    combined = combined.reindex(columns=[c for c in COLS_OUT_ORDER if c in combined.columns])

    return loan, cancel, combined

def _sort_for_date_sheet(df_date: pd.DataFrame, config: Dict[str, Any]) -> pd.DataFrame:
    """
    日付シートの並べ替えを設定に基づいて実行します。
    """
    staff_order = config.get("staff_order", [])
    order_dict = {name: i for i, name in enumerate(staff_order)}
    sort_keys_config = config.get("sort_keys", [
        {"key": "識別", "ascending": False},
        {"key": "担当", "ascending": True},
        {"key": "時間", "ascending": True}
    ])

    temp_cols = []
    sort_by_cols = []
    ascending_list = []

    for i, item in enumerate(sort_keys_config):
        col_name = item.get("key", "")
        asc = item.get("ascending", True)
        temp_col = f"_sort_key_{i}"

        if col_name == "担当":
            # 担当順カスタムソート（未掲載は末尾）
            s_val = df_date['担当'].astype(str).map(order_dict).fillna(99999).astype(int)
        elif col_name == "時間":
            # 時間パース
            s_val = _parse_time_like(df_date.get('時間', pd.Series()))
        elif col_name == "識別":
            # 識別（文字列ソート）
            s_val = df_date.get('識別', pd.Series()).astype(str)
        else:
            if col_name in df_date.columns:
                s_val = df_date[col_name].astype(str)
            else:
                s_val = pd.Series([""] * len(df_date), index=df_date.index)

        df_date = df_date.assign(**{temp_col: s_val})
        temp_cols.append(temp_col)
        sort_by_cols.append(temp_col)
        ascending_list.append(asc)

    if sort_by_cols:
        df_sorted = df_date.sort_values(by=sort_by_cols, ascending=ascending_list).drop(columns=temp_cols)
    else:
        df_sorted = df_date

    return df_sorted

def _apply_excel_styles(out_path: Path, config: Dict[str, Any], log_cb: Optional[Callable[[str], None]] = None):
    """openpyxlを用いて罫線、列幅、行高、配置スタイルを適用します"""
    if log_cb:
        log_cb("Excelの書式設定を適用中...")

    formatting = config.get("formatting", {})
    row_height = formatting.get("row_height", 25)
    width_rules = formatting.get("col_widths", {})
    default_width = formatting.get("default_col_width", 11.43)

    wb = load_workbook(out_path)

    dotted = Side(style='dotted')
    border_all = Border(left=dotted, right=dotted, top=dotted, bottom=dotted)
    align = Alignment(horizontal='center', vertical='center', shrink_to_fit=True, wrap_text=False)
    header_font = Font(bold=True)

    for ws in wb.worksheets:
        max_row = ws.max_row
        max_col = ws.max_column

        # 罫線・高さ・配置の設定
        for r in range(1, max_row + 1):
            ws.row_dimensions[r].height = row_height
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.border = border_all
                if r == 1:
                    cell.font = header_font
                cell.alignment = align

        # 列幅の設定
        for c in range(1, max_col + 1):
            col_letter = ws.cell(row=1, column=c).column_letter
            header_val = ws.cell(row=1, column=c).value
            width = width_rules.get(header_val, default_width)
            ws.column_dimensions[col_letter].width = width

    wb.save(out_path)

def process_excel(input_path: str, output_path: Optional[str] = None, config: Optional[Dict[str, Any]] = None, log_cb: Optional[Callable[[str], None]] = None) -> str:
    """
    Excelファイルのソートと書式設定を実行します。
    """
    src = Path(input_path)
    if not src.exists():
        raise FileNotFoundError(f"入力ファイルが見つかりません: {src}")

    if config is None:
        config = load_config()

    if output_path is None:
        # デフォルト出力パス：同じディレクトリに [元ファイル名]_sorted.xlsx
        output_path = str(src.parent / f"{src.stem}_sorted.xlsx")
    
    out_file = Path(output_path)

    # 1. 読み込み
    df = _read_sheet(src, log_cb=log_cb)

    # 2. データ整形
    if log_cb:
        log_cb("データを並べ替え/集計中...")
    loan, cancel, combined = _make_tables(df)

    sheet_opts = config.get("sheet_options", {})
    
    # 3. Excel出力
    if log_cb:
        log_cb("シートを出力中...")
    with pd.ExcelWriter(out_file, engine='openpyxl') as writer:
        if sheet_opts.get("loan_sheet", True):
            loan.to_excel(writer, sheet_name='sorted_by_融資実行日', index=False)
        if sheet_opts.get("cancel_sheet", True):
            cancel.to_excel(writer, sheet_name='sorted_by_金消日・面談日', index=False)
        if sheet_opts.get("combined_sheet", True):
            combined.to_excel(writer, sheet_name='統合データ', index=False)

        # 日付ごとのシート
        if sheet_opts.get("date_sheets", True) and '日付' in combined.columns:
            unique_dates = combined['日付'].dropna().unique()
            if log_cb:
                log_cb(f"日付別シート作成中 (計 {len(unique_dates)} 日分)...")
            for date in unique_dates:
                df_date = combined[combined['日付'] == date].copy()
                df_date = _sort_for_date_sheet(df_date, config)
                sheet_name = str(date).replace('/', '-')
                # Excelシート名制限（31文字以下）
                sheet_name = sheet_name[:31]
                df_date.to_excel(writer, sheet_name=sheet_name, index=False)

    # 4. 書式適用
    _apply_excel_styles(out_file, config, log_cb=log_cb)

    if log_cb:
        log_cb(f"処理が完了しました: {out_file.name}")

    return str(out_file.resolve())
